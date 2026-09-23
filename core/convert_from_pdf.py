"""Convert PDFs to other formats."""
from pathlib import Path

from .utils import (
    OUTPUT_DIR, logger,
    print_success, print_error,
    ask_path, ask_save_path,
)


# ------------------------------------------------------------------
# PDF → Word (pdf2docx)
# ------------------------------------------------------------------
def pdf_to_word():
    print("\n--- PDF → Word ---")
    src = ask_path("PDF file: ", file_ext=(".pdf",))
    dst = ask_save_path(src.stem, ".docx")
    try:
        from pdf2docx import Converter
        cv = Converter(str(src))
        cv.convert(str(dst))
        cv.close()
        print_success(f"Saved: {dst}")
    except Exception as e:
        print_error(f"PDF to Word failed: {e}")


# ------------------------------------------------------------------
# PDF → Images
# ------------------------------------------------------------------
def pdf_to_images():
    print("\n--- PDF → Images (JPG/PNG) ---")
    src = ask_path("PDF file: ", file_ext=(".pdf",))
    fmt = input("Format (jpg/png) [jpg]: ").strip().lower() or "jpg"
    fmt = "jpeg" if fmt == "jpg" else fmt
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(src)
        out_dir = OUTPUT_DIR / f"{src.stem}_images"
        out_dir.mkdir(parents=True, exist_ok=True)
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=150)
            pix.save(out_dir / f"page_{i+1:03d}.{fmt}")
        doc.close()
        print_success(f"Extracted {i+1} page(s) to {out_dir}")
    except Exception as e:
        print_error(f"PDF to images failed: {e}")


# ------------------------------------------------------------------
# PDF → Text
# ------------------------------------------------------------------
def pdf_to_text():
    print("\n--- PDF → Text ---")
    src = ask_path("PDF file: ", file_ext=(".pdf",))
    dst = ask_save_path(src.stem, ".txt")
    try:
        import fitz
        doc = fitz.open(src)
        text = "\n\n".join(page.get_text() for page in doc)
        doc.close()
        dst.write_text(text, encoding="utf-8")
        print_success(f"Saved: {dst}")
    except Exception as e:
        print_error(f"PDF to text failed: {e}")


# ------------------------------------------------------------------
# Extract embedded images
# ------------------------------------------------------------------
def extract_images():
    print("\n--- Extract Embedded Images from PDF ---")
    src = ask_path("PDF file: ", file_ext=(".pdf",))
    try:
        import fitz
        doc = fitz.open(src)
        out_dir = OUTPUT_DIR / f"{src.stem}_images"
        out_dir.mkdir(parents=True, exist_ok=True)
        count = 0
        for pno in range(len(doc)):
            for img_index, img in enumerate(doc.get_page_images(pno, full=True)):
                xref = img[0]
                base = doc.extract_image(xref)
                ext = base["ext"]
                data = base["image"]
                (out_dir / f"p{pno+1}_img{img_index+1}.{ext}").write_bytes(data)
                count += 1
        doc.close()
        print_success(f"Extracted {count} image(s) to {out_dir}")
    except Exception as e:
        print_error(f"Image extraction failed: {e}")


# ------------------------------------------------------------------
# PDF → Excel (basic text dump using openpyxl)
# ------------------------------------------------------------------
def pdf_to_excel():
    print("\n--- PDF → Excel (text rows) ---")
    src = ask_path("PDF file: ", file_ext=(".pdf",))
    dst = ask_save_path(src.stem, ".xlsx")
    try:
        import fitz
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "PDF Text"
        doc = fitz.open(src)
        row = 1
        for pno, page in enumerate(doc):
            ws.cell(row=row, column=1, value=f"--- Page {pno+1} ---")
            row += 1
            for line in page.get_text().splitlines():
                ws.cell(row=row, column=1, value=line)
                row += 1
        doc.close()
        wb.save(dst)
        print_success(f"Saved: {dst}")
    except Exception as e:
        print_error(f"PDF to Excel failed: {e}")


# ------------------------------------------------------------------
# PDF → PowerPoint (render each page as an image slide)
# ------------------------------------------------------------------
def pdf_to_powerpoint():
    print("\n--- PDF → PowerPoint (image slides) ---")
    src = ask_path("PDF file: ", file_ext=(".pdf",))
    dst = ask_save_path(src.stem, ".pptx")
    try:
        import fitz
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image
        import io

        doc = fitz.open(src)
        prs = Presentation()
        blank = prs.slide_layouts[6]

        for page in doc:
            pix = page.get_pixmap(dpi=150)
            img_bytes = pix.tobytes("png")
            img = Image.open(io.BytesIO(img_bytes))
            slide = prs.slides.add_slide(blank)
            tmp = OUTPUT_DIR / "_tmp_page.png"
            img.save(tmp)
            slide.shapes.add_picture(str(tmp), Inches(0), Inches(0),
                                     width=prs.slide_width)
            tmp.unlink(missing_ok=True)
        doc.close()
        prs.save(dst)
        print_success(f"Saved: {dst}")
    except Exception as e:
        print_error(f"PDF to PowerPoint failed: {e}")


# ------------------------------------------------------------------
# Dispatch
# ------------------------------------------------------------------
FROM_PDF_MENU = {
    "1": ("PDF to Word (.docx)", pdf_to_word),
    "2": ("PDF to PowerPoint (.pptx)", pdf_to_powerpoint),
    "3": ("PDF to Excel (.xlsx)", pdf_to_excel),
    "4": ("PDF to Images (JPG/PNG)", pdf_to_images),
    "5": ("PDF to Text (.txt)", pdf_to_text),
    "6": ("Extract Embedded Images", extract_images),
}


def run():
    while True:
        print("\n--- Convert FROM PDF ---")
        for k, (label, _) in FROM_PDF_MENU.items():
            print(f" {k:>2}. {label}")
        print("  0. Back")
        choice = input("Choice: ").strip()
        if choice == "0":
            return
        if choice in FROM_PDF_MENU:
            FROM_PDF_MENU[choice][1]()
        else:
            print("❌ Invalid choice.")
