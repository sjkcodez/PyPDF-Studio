"""AST → PPTX writer.

Each Heading(level=1) starts a new slide.
Content nodes under it form the slide body.
"""
from pptx import Presentation
from pptx.util import Inches
from core.ast import (
    Document, Heading, Paragraph, BulletList, NumberedList, Table,
)


def _split_slides(ast: Document):
    slides = []
    current = None

    for node in ast.children:
        if isinstance(node, Heading) and node.level == 1:
            if current is not None:
                slides.append(current)
            title = node.text
            if title.lower().startswith("slide") and ":" in title:
                title = title.split(":", 1)[1].strip()
            current = {"title": title or "Slide", "body": []}
        else:
            if current is None:
                current = {
                    "title": ast.metadata.get("title", "Slide"),
                    "body": [],
                }
            current["body"].append(node)

    if current is not None:
        slides.append(current)

    return slides


def write(ast: Document, output_path: str):
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[1]  # Title and Content

    for slide_data in _split_slides(ast):
        slide = prs.slides.add_slide(layout)

        if slide.shapes.title:
            slide.shapes.title.text = slide_data["title"]

        body_ph = None
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1:
                body_ph = ph
                break
        if body_ph is None:
            continue

        tf = body_ph.text_frame
        first = True

        for node in slide_data["body"]:
            if isinstance(node, BulletList):
                items, level = node.items, 0
            elif isinstance(node, NumberedList):
                items, level = node.items, 0
            elif isinstance(node, Paragraph):
                items, level = [node.text], 0
            elif isinstance(node, Table):
                items = [" | ".join(node.headers)] + [" | ".join(r) for r in node.rows]
                level = 0
            else:
                continue

            for item in items:
                if first:
                    tf.text = item
                    first = False
                else:
                    p = tf.add_paragraph()
                    p.text = item
                    p.level = level

    prs.save(output_path)
