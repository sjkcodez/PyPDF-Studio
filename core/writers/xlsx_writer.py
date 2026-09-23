"""AST → XLSX writer.

Each Heading(level=1) starts a new sheet.
Each Table node fills that sheet with rows.
"""
from openpyxl import Workbook
from core.ast import Document, Heading, Table


def write(ast: Document, output_path: str):
    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    current_ws = None

    def ensure_sheet(name="Sheet"):
        nonlocal current_ws
        name = (name or "Sheet").strip()[:31] or "Sheet"
        base, i = name, 2
        while name in wb.sheetnames:
            suffix = f"_{i}"
            name = base[:31 - len(suffix)] + suffix
            i += 1
        current_ws = wb.create_sheet(title=name)
        return current_ws

    for node in ast.children:
        if isinstance(node, Heading) and node.level == 1:
            title = node.text
            if title.lower().startswith("sheet:"):
                title = title.split(":", 1)[1].strip()
            ensure_sheet(title or "Sheet")
        elif isinstance(node, Table):
            if current_ws is None:
                ensure_sheet("Sheet")
            current_ws.append(node.headers)
            for row in node.rows:
                current_ws.append(row)

    if not wb.sheetnames:
        wb.create_sheet(title="Sheet")

    wb.save(output_path)
