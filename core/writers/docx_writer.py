"""AST → DOCX writer."""
from docx import Document as DocxDocument
from core.ast import (
    Document, Heading, Paragraph, BulletList, NumberedList,
    CodeBlock, BlockQuote, Table,
)


def write(ast: Document, output_path: str):
    doc = DocxDocument()

    if ast.metadata.get("title"):
        doc.add_heading(ast.metadata["title"], level=0)

    for node in ast.children:
        if isinstance(node, Heading):
            doc.add_heading(node.text, level=min(node.level, 9))
        elif isinstance(node, Paragraph):
            doc.add_paragraph(node.text)
        elif isinstance(node, BulletList):
            for item in node.items:
                doc.add_paragraph(item, style="List Bullet")
        elif isinstance(node, NumberedList):
            for item in node.items:
                doc.add_paragraph(item, style="List Number")
        elif isinstance(node, CodeBlock):
            p = doc.add_paragraph()
            run = p.add_run(node.code)
            run.font.name = "Courier New"
        elif isinstance(node, BlockQuote):
            doc.add_paragraph(node.text, style="Intense Quote")
        elif isinstance(node, Table):
            table = doc.add_table(rows=1, cols=len(node.headers))
            table.style = "Light Grid Accent 1"
            for i, h in enumerate(node.headers):
                table.rows[0].cells[i].text = h
            for row in node.rows:
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    if i < len(cells):
                        cells[i].text = val

    doc.save(output_path)
