"""PPTX → AST reader.

Each slide becomes:
    Heading(level=1, "Slide N: <title>")
    Paragraph / BulletList nodes for the content
"""
from pathlib import Path
from pptx import Presentation
from core.ast import Document, Heading, Paragraph, BulletList


def read(path: str) -> Document:
    prs = Presentation(path)
    ast = Document()
    ast.metadata["title"] = Path(path).stem

    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title_text = title_shape.text_frame.text.strip() if title_shape else ""

        ast.children.append(
            Heading(level=1, text=f"Slide {i}: {title_text or '(untitled)'}")
        )

        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if not shape.has_text_frame:
                continue

            texts = [
                p.text.strip()
                for p in shape.text_frame.paragraphs
                if p.text and p.text.strip()
            ]
            if not texts:
                continue

            if len(texts) == 1:
                ast.children.append(Paragraph(text=texts[0]))
            else:
                ast.children.append(BulletList(items=texts))

    return ast
