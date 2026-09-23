"""DOCX → AST reader."""
from docx import Document as DocxDocument
from core.ast import (
    Document, Heading, Paragraph, BulletList,
    NumberedList, Table,
)


def _style_level(style_name: str):
    """Return heading level from style name, else None."""
    s = style_name.lower()
    if s.startswith("heading"):
        digits = "".join(ch for ch in s if ch.isdigit())
        return int(digits) if digits else 1
    if s == "title":
        return 1
    return None


def read(path: str) -> Document:
    doc = DocxDocument(path)
    ast = Document()
    ast.metadata["title"] = doc.core_properties.title or ""

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        level = _style_level(para.style.name if para.style else "")
        if level is not None:
            ast.children.append(Heading(level=level, text=text))
            continue

        style_name = (para.style.name if para.style else "").lower()
        if "list bullet" in style_name:
            ast.children.append(BulletList(items=[text]))
        elif "list number" in style_name:
            ast.children.append(NumberedList(items=[text]))
        else:
            ast.children.append(Paragraph(text=text))

    for tbl in doc.tables:
        if not tbl.rows:
            continue
        headers = [c.text.strip() for c in tbl.rows[0].cells]
        rows = [[c.text.strip() for c in r.cells] for r in tbl.rows[1:]]
        ast.children.append(Table(headers=headers, rows=rows))

    return ast
